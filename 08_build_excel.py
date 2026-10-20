"""Build the stakeholder Excel workbook: a shrinkage calculator plus accuracy and anomaly sheets.

Every number a reader might change or audit is a live Excel formula, not a pasted value.
openpyxl writes formulas but does not calculate them; Excel calculates on open.
Run 07_export_powerbi.py first: this reads its tables from data/powerbi/.
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

IN_DIR = Path("data/powerbi")
OUT_PATH = Path("reports/shrinkage-calculator.xlsx")

TITLE = Font(size=14, bold=True)
HEADER = Font(bold=True)
NOTE = Font(italic=True, color="595959")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HOT_FILL = PatternFill("solid", fgColor="F8CBAD")
COLD_FILL = PatternFill("solid", fgColor="BDD7EE")
PCT = "0.0%"
RATE = "0.0000"


def add_table(ws, name, df, first_row=1):
    """Write df at first_row as a formatted, filterable Excel Table. Returns the last row."""
    for j, col in enumerate(df.columns, start=1):
        ws.cell(first_row, j, col).font = HEADER
    for i, row in enumerate(df.itertuples(index=False), start=first_row + 1):
        for j, value in enumerate(row, start=1):
            ws.cell(i, j, value)
    last_row = first_row + len(df)
    ref = f"A{first_row}:{get_column_letter(len(df.columns))}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    return last_row


def set_widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def write_priors(wb, dim_season):
    """Season priors as an Excel Table; the calculator looks alpha and beta up here."""
    ws = wb.create_sheet("Priors")
    cols = ["SEASON_YEAR", "N_PLAYERS", "ALPHA", "BETA", "PRIOR_MEAN", "STRENGTH"]
    last = add_table(ws, "Priors", dim_season[cols])
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=3, max_col=6):
        for cell in row:
            cell.number_format = "0.0000" if cell.column == 5 else "0.00"
    set_widths(ws, {"A": 14, "B": 12, "C": 10, "D": 10, "E": 13, "F": 11})
    return last


def write_calculator(wb, priors_last_row):
    """Inputs (yellow) -> season prior lookup -> shrunk estimate and credible interval."""
    ws = wb.active
    ws.title = "Calculator"
    seasons = f"Priors!$A$2:$A${priors_last_row}"

    def lookup(col):
        return f"=INDEX(Priors!${col}$2:${col}${priors_last_row},MATCH($B$5,{seasons},0))"

    rows = [
        # row, label, value or formula, number format, note
        (5, "Season", "2024-25", None, "Which season's league prior to shrink toward"),
        (6, "Makes (3PM)", 34, "0", "Threes made"),
        (7, "Attempts (3PA)", 76, "0", "Threes attempted"),
        (8, "Credible level", 0.9, "0%", "Width of the interval below"),
        (11, "Alpha", lookup("C"), "0.00", "Fitted Beta prior: league 'phantom makes'"),
        (12, "Beta", lookup("D"), "0.00", "Fitted Beta prior: league 'phantom misses'"),
        (13, "Prior mean", "=B11/(B11+B12)", PCT, "League-average 3P% that season"),
        (14, "Phantom attempts", "=B11+B12", "0", "Attempts needed before the player's own data gets half the weight"),
        (17, "Raw 3P%", "=IF(B7>0,B6/B7,NA())", PCT, "What the box score says"),
        (18, "Shrunk 3P%", "=(B6+B11)/(B7+B14)", PCT, "Posterior mean: (makes + alpha) / (attempts + alpha + beta)"),
        (19, "Weight on player's data", "=B7/(B7+B14)", PCT, "Shrunk = weight x raw + (1 - weight) x prior mean"),
        (20, "Shift", "=B18-B17", "+0.0%;-0.0%", "How far shrinkage moved the estimate"),
        (21, "Interval low", "=_xlfn.BETA.INV((1-B8)/2,B11+B6,B12+B7-B6)", PCT, "Posterior Beta quantile"),
        (22, "Interval high", "=_xlfn.BETA.INV(1-(1-B8)/2,B11+B6,B12+B7-B6)", PCT, "Posterior Beta quantile"),
    ]

    ws["A1"] = "3P% Shrinkage Calculator"
    ws["A1"].font = TITLE
    ws["A2"] = "Change the yellow cells. Everything else is a formula."
    ws["A2"].font = NOTE
    for row, text in [(4, "Inputs"), (10, "Season prior"), (16, "Estimate")]:
        ws.cell(row, 1, text).font = HEADER

    for row, label, value, fmt, note in rows:
        ws.cell(row, 1, label)
        cell = ws.cell(row, 2, value)
        if fmt:
            cell.number_format = fmt
        ws.cell(row, 3, note).font = NOTE
        if row <= 8:
            cell.fill = INPUT_FILL
    ws["B18"].font = HEADER

    season_dv = DataValidation(type="list", formula1=seasons, allow_blank=False)
    attempts_dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1",
                                 error="Attempts must be a whole number of at least 1.")
    makes_dv = DataValidation(type="custom", formula1="AND(INT(B6)=B6,B6>=0,B6<=B7)",
                              error="Makes must be a whole number between 0 and attempts.")
    level_dv = DataValidation(type="decimal", operator="between", formula1="0.5", formula2="0.99",
                              error="Use a level between 50% and 99%.")
    for dv, cell in [(season_dv, "B5"), (makes_dv, "B6"), (attempts_dv, "B7"), (level_dv, "B8")]:
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(cell)

    ws.conditional_formatting.add("B20", CellIsRule(operator="lessThan", formula=["-0.02"], fill=COLD_FILL))
    ws.conditional_formatting.add("B20", CellIsRule(operator="greaterThan", formula=["0.02"], fill=HOT_FILL))
    set_widths(ws, {"A": 24, "B": 12, "C": 70})


def write_data(wb, fact_forecast, fact_pair, dim_method):
    """Forecast errors in long format: the source for the Accuracy formulas."""
    ws = wb.create_sheet("Data")
    df = (fact_forecast
          .merge(fact_pair[["PAIR_ID", "PAIR"]], on="PAIR_ID")
          .merge(dim_method[["METHOD_KEY", "METHOD_LABEL"]], on="METHOD_KEY"))
    cols = ["PAIR_ID", "PAIR", "VOLUME_BUCKET", "IN_ML_WINDOW", "METHOD_LABEL",
            "PREDICTION", "ACTUAL", "SQ_ERROR"]
    add_table(ws, "ForecastData", df[cols])
    set_widths(ws, {"B": 22, "C": 15, "D": 14, "E": 18, "F": 12, "G": 10, "H": 12})
    ws.freeze_panes = "A2"


def accuracy_grid(ws, top, title, labels, columns, match_col):
    """RMSE grid: one row per method, one column per value of match_col, plus pooled.

    Every cell is SQRT(AVERAGEIFS(squared error, ...)) over the Data sheet, restricted to
    the ML test window so all five methods are scored on the same player pairs.
    """
    ws.cell(top, 1, title).font = HEADER
    header = top + 1
    ws.cell(header, 1, "Method").font = HEADER
    for j, value in enumerate(columns, start=2):
        ws.cell(header, j, value).font = HEADER
    pooled_col = len(columns) + 2
    gain_col = pooled_col + 1
    ws.cell(header, pooled_col, "Pooled").font = HEADER
    ws.cell(header, gain_col, "RMSE cut vs. EB").font = HEADER

    sq, method, window = "Data!$H:$H", "Data!$E:$E", "Data!$D:$D"
    eb_row = header + 1 + labels.index("EB shrunk")
    for i, label in enumerate(labels, start=header + 1):
        ws.cell(i, 1, label)
        for j in range(2, pooled_col):
            col = get_column_letter(j)
            ws.cell(i, j, f"=SQRT(AVERAGEIFS({sq},{method},$A{i},{match_col},{col}${header},{window},1))")
        ws.cell(i, pooled_col, f"=SQRT(AVERAGEIFS({sq},{method},$A{i},{window},1))")
        p = get_column_letter(pooled_col)
        ws.cell(i, gain_col, f"=1-{p}{i}/{p}${eb_row}")
        for j in range(2, pooled_col + 1):
            ws.cell(i, j).number_format = RATE
        ws.cell(i, gain_col).number_format = "+0.0%;-0.0%;0.0%"
    return header + len(labels)


def write_accuracy(wb, fact_pair, dim_method):
    """RMSE by method x test season and method x volume bucket, as formulas."""
    ws = wb.create_sheet("Accuracy", 1)
    ws["A1"] = "Forecast accuracy: next-season 3P%"
    ws["A1"].font = TITLE
    ws["A2"] = ("RMSE over the ML test window (1,243 player pairs), so every method is "
                "scored on the same players. Lower is better.")
    ws["A2"].font = NOTE

    labels = dim_method.sort_values("SORT_ORDER")["METHOD_LABEL"].tolist()
    test_pairs = sorted(fact_pair.loc[fact_pair["IN_ML_WINDOW"] == 1, "PAIR"].unique())
    end = accuracy_grid(ws, 4, "By test season", labels, test_pairs, "Data!$B:$B")
    buckets = ["<100 3PA", "100-199", "200-399", "400+"]
    accuracy_grid(ws, end + 3, "By season-t volume", labels, buckets, "Data!$C:$C")
    set_widths(ws, {"A": 20, **{get_column_letter(j): 20 for j in range(2, 9)}})


def write_anomalies(wb, fact_pair):
    """Flagged player pairs, most surprising first."""
    ws = wb.create_sheet("Anomalies", 2)
    flagged = fact_pair[fact_pair["IS_ANOMALY"] == 1].copy()
    flagged["SURPRISE"] = (flagged["PIT"] - 0.5).abs()
    flagged = flagged.sort_values("SURPRISE", ascending=False)

    cols = {
        "PLAYER_NAME": "Player", "PAIR": "Seasons", "FG3M": "3PM t", "FG3A": "3PA t",
        "FG3_PCT_SHRUNK": "Shrunk 3P% t", "FG3M_NEXT": "3PM t+1", "FG3A_NEXT": "3PA t+1",
        "ACTUAL": "3P% t+1", "PI_LOW": "Interval low (3PM)", "PI_HIGH": "Interval high (3PM)",
        "PIT": "Percentile", "DIRECTION": "Direction",
    }
    out = flagged[list(cols)].rename(columns=cols)
    out[["Interval low (3PM)", "Interval high (3PM)"]] = out[["Interval low (3PM)", "Interval high (3PM)"]].astype(int)

    ws["A1"] = f"Anomalies: {len(out)} of {len(fact_pair):,} player pairs ({len(out) / len(fact_pair):.1%})"
    ws["A1"].font = TITLE
    ws["A2"] = ("Next-season makes fell outside the 95% Beta-Binomial predictive interval. "
                "About 4.2% would be flagged by chance alone if the model were exactly right.")
    ws["A2"].font = NOTE
    last = add_table(ws, "Anomalies", out, first_row=4)

    for row in ws.iter_rows(min_row=5, max_row=last):
        row[4].number_format = PCT
        row[7].number_format = PCT
        row[10].number_format = "0.000"
    rng = f"A5:L{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$L5="hot"'], fill=HOT_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$L5="cold"'], fill=COLD_FILL))
    set_widths(ws, {"A": 24, "B": 22, **{c: 13 for c in "CDEFGHIJKL"}})
    ws.freeze_panes = "B5"
    for cell in ws[4]:
        cell.alignment = Alignment(wrap_text=True)


def main():
    fact_pair = pd.read_csv(IN_DIR / "fact_pair.csv")
    fact_forecast = pd.read_csv(IN_DIR / "fact_forecast.csv")
    dim_season = pd.read_csv(IN_DIR / "dim_season.csv")
    dim_method = pd.read_csv(IN_DIR / "dim_method.csv")

    wb = Workbook()
    priors_last = write_priors(wb, dim_season)
    write_calculator(wb, priors_last)
    write_data(wb, fact_forecast, fact_pair, dim_method)
    write_accuracy(wb, fact_pair, dim_method)
    write_anomalies(wb, fact_pair)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"saved -> {OUT_PATH}  (sheets: {', '.join(wb.sheetnames)})")


if __name__ == "__main__":
    main()
